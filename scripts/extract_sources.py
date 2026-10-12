#!/usr/bin/env python3
"""Read-only extraction of Word Flat OPC and master indexing workbook.
All joins keep their source locator. No source cells/tokens are repaired.
"""
from pathlib import Path
from collections import Counter,defaultdict
from copy import deepcopy
from lxml import etree as E
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as COL
import re,json,hashlib,datetime,argparse

P=Path(__file__).resolve().parents[1]
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
PKG='http://schemas.microsoft.com/office/2006/xmlPackage'
R='http://schemas.openxmlformats.org/package/2006/relationships'
NS={'w':W,'pkg':PKG,'r':R}
A='A228401'; TITLE='כי תצא למלחמה'

def val(x):
 if isinstance(x,(datetime.datetime,datetime.date,datetime.time)):return x.isoformat()
 return x

def txt(x):return '' if x is None else str(x)
def get(v,i):return v[i] if i<len(v) else None
def norm(x):return re.sub(r'\s+',' ',txt(x)).strip()
def ids(x):return re.findall(r'A\d+(?:-[A-Z]+\d+(?:\.\w+)*)?|per\d+',txt(x))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def serial(x):
 if isinstance(x,(datetime.datetime,datetime.date,datetime.time)):return x.isoformat()
 if hasattr(x,'text') and hasattr(x,'ref'):return {'formula':x.text,'ref':x.ref,'type':type(x).__name__}
 raise TypeError('Unsupported '+type(x).__name__)
def jwrite(p,d):p.write_text(json.dumps(d,ensure_ascii=False,default=serial,indent=2),encoding='utf8')

def run(xmlpath,xlsxpath):
 wb=load_workbook(xlsxpath,data_only=True); wf=load_workbook(xlsxpath,data_only=False,read_only=True)
 sheets=[]; rows_by={}; issues=[]
 def issue(id,severity,title,detail,source,impact):
  issues.append(dict(id=id,severity=severity,title=title,detail=detail,source=source,impact=impact))
 for ws in wb:
  entries=[];headers=[];errs=[]; formula_count=0
  for rr in ws:
   values=[val(c.value) for c in rr]
   if not any(v is not None and v!='' for v in values):continue
   while values and values[-1] is None:values.pop()
   if not headers:
    headers=values;continue
   cells={f'{COL(i+1)}{rr[0].row}':v for i,v in enumerate(values) if v is not None}
   rec={'sheet':ws.title,'row':rr[0].row,'values':values,'cells':cells}
   href={c.coordinate:c.hyperlink.target for c in rr if c.hyperlink and c.hyperlink.target}
   if href:rec['links']=href
   coms={c.coordinate:c.comment.text for c in rr if c.comment}
   if coms:rec['comments']=coms
   direct=any(re.search(r'A228401(?!\d)',txt(v)) or norm(v)==TITLE for v in values)
   rec['direct_match']=direct
   errors=[c.coordinate for c in rr if c.data_type=='e' or (isinstance(c.value,str) and c.value in ('#REF!','#N/A','#VALUE!','#DIV/0!','#NAME?'))]
   if errors:errs.extend([{'cell':c,'value':cells.get(c),'target':direct} for c in errors])
   entries.append(rec)
  # Do not expand styled blank tail into a dataset.
  fws=wf[ws.title]
  forms=[]
  for rr in fws.iter_rows(max_row=max([r['row'] for r in entries],default=1),max_col=max((len(r['values']) for r in entries),default=1)):
   for c in rr:
    if c.data_type=='f': forms.append({'cell':c.coordinate,'formula':c.value})
  rows_by[ws.title]=entries
  sheets.append(dict(name=ws.title,state=ws.sheet_state,dimension=ws.calculate_dimension(),headers=headers,nonempty_rows=len(entries),direct_rows=sum(r['direct_match'] for r in entries),formulas=forms,errors=errs,rows=entries))
 def choose(prefix):return next(s for s in sheets if s['name'].startswith(prefix))
 def matches(prefix):return [r for r in choose(prefix)['rows'] if r['direct_match']]

 # Canonical sheet identities, not guessed relational table names.
 prod=matches('A -')[0]
 scene_records={r['values'][1]:{'id':r['values'][1],'title':get(r['values'],3),'note':get(r['values'],4),'record':r} for r in matches('S -')}
 text_records={r['values'][0]:{'id':r['values'][0],'title':get(r['values'],4),'target':get(r['values'],5),'genre':get(r['values'],6),'record':r} for r in matches('T -')}
 progs={r['values'][3]:{'id':r['values'][3],'target':get(r['values'],4),'title':get(r['values'],13),'order':get(r['values'],14),'parent':r['values'][3].rsplit('.PS',1)[0],'parent_kind':'derived_from_identifier','record':r} for r in matches('PS -')}
 graphics={r['values'][1]:{'id':r['values'][1],'kind':get(r['values'],5),'note':get(r['values'],6),'logical_page':get(r['values'],3),'image_page':get(r['values'],4),'record':r} for r in matches('P -')}
 roles={r['values'][3]:{'id':r['values'][3],'s':get(r['values'],2),'character':get(r['values'],4),'actor':get(r['values'],6),'person':get(r['values'],7),'note':get(r['values'],8),'record':r} for r in matches('תפקידים מקושרים')}
 all_persons={r['values'][1]:{'id':r['values'][1],'name':get(r['values'],0),'first':get(r['values'],2),'last':get(r['values'],3),'wikidata':get(r['values'],4),'viaf':get(r['values'],5),'record':r} for r in choose('ישויות')['rows'] if re.fullmatch(r'per\d+',txt(get(r['values'],1)))}
 p_by_name=defaultdict(list)
 for pid,p in all_persons.items():p_by_name[norm(p['name'])].append(pid)
 used_persons=set();name_links=[];person_errors=[]
 for rid,r in roles.items():
  if r['person'] in all_persons:used_persons.add(r['person'])
  elif r['person'] not in (None,'','#N/A'):person_errors.append((rid,r['person']))
 for rec in [prod]+[x['record'] for x in progs.values()]+matches('תפקידים בהצגה'):
  for ci,v in enumerate(rec['values']):
   if norm(v) in p_by_name and len(p_by_name[norm(v)])==1:
    pid=p_by_name[norm(v)][0];used_persons.add(pid)
    name_links.append({'person':pid,'source':f"{rec['sheet']}!{COL(ci+1)}{rec['row']}",'name':v,'kind':'unique_exact_name_match_not_id'})
 for rid,pid in person_errors:issue('unresolved-person-'+rid,'warning','קוד אדם שאינו נפתר',pid,roles[rid]['record'],'אין לייצר אדם חלופי על סמך השם בלבד.')
 media=[];quarantine=[]
 for rr in matches('אודיו'):
  v=rr['values'];rec={'id':f'av-row-{rr["row"]}','production':get(v,0),'production_title':get(v,1),'title':get(v,2),'target':get(v,3),'kind':get(v,4),'archive':get(v,5),'url':get(v,6),'date':get(v,7),'extra':v[8:],'record':rr}
  if get(v,0)==A and norm(get(v,1))!=TITLE and not txt(get(v,3)).startswith(A):
   quarantine.append(rec)
   issue('av-conflict-'+str(rr['row']),'error','שיוך סותר של חומר שמע',f"A={get(v,0)}; שם={get(v,1)}; S={get(v,3)}",rr,'הרשומה נשמרת בדוח ואינה מצורפת אוטומטית להצגה.')
  else:media.append(rec)

 # Word Flat OPC package: XML is a container, not TEI.
 root=E.parse(str(xmlpath),E.XMLParser(resolve_entities=False,no_network=True))
 parts={x.get('{'+PKG+'}name'):x for x in root.xpath('/pkg:package/pkg:part',namespaces=NS)}
 doc=parts['/word/document.xml'].find('pkg:xmlData',NS)[0]
 styles={}
 if '/word/styles.xml' in parts:
  for s in parts['/word/styles.xml'].xpath('.//w:style',namespaces=NS):styles[s.get('{'+W+'}styleId')]=s
 def attrs_rpr(rpr):
  if rpr is None:return {}
  return {E.QName(x).localname:dict(x.attrib) for x in rpr}
 def props(simple):
  return {k:v.get('{'+W+'}val',True) for k,v in simple.items()}
 def rtext(r):
  v=[]
  for x in r:
   ln=E.QName(x).localname
   if ln in ('t','delText'):v.append(x.text or '')
   elif ln=='tab':v.append('\t')
   elif ln in ('br','cr'):v.append('\n')
   elif ln=='noBreakHyphen':v.append('\u2011')
  return ''.join(v)
 def relationships(part):
  if part not in parts:return {}
  return {r.get('Id'):{'target':r.get('Target'),'type':r.get('Type'),'mode':r.get('TargetMode')} for r in parts[part].xpath('.//r:Relationship',namespaces=NS)}
 links=relationships('/word/_rels/document.xml.rels');flinks=relationships('/word/_rels/footnotes.xml.rels')
 footnotes={}
 f_part=parts.get('/word/footnotes.xml')
 if f_part is not None:
  for f in f_part.xpath('.//w:footnote',namespaces=NS):
   fid=f.get('{'+W+'}id')
   if int(fid)<=0:continue
   content='\n'.join(''.join(p.xpath('.//w:t/text()|.//w:delText/text()',namespaces=NS)) for p in f.xpath('./w:p',namespaces=NS)).strip()
   hrefs=[flinks.get(h.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'),{}) for h in f.xpath('.//w:hyperlink',namespaces=NS)]
   cats=[]
   for c in ['תקשורת המונים','אישים','מקום','נשק','מושג','מוסד','שיר','אירוע','אקטואליה','תרבות','פארודיה']:
    if re.search(r'אזכור\s+'+re.escape(c),content):cats.append(c)
   footnotes[fid]={'id':fid,'text':content,'links':hrefs,'categories':cats,'internal':content.lstrip().startswith('<פנימית') or content.startswith('פנימית')}
 paragraphs=[];raw_runs_count=0;footrefs=[];colors=Counter();highlights=Counter();sems=Counter();offset_total=0
 role_pat=re.compile(r'<(A228401-R\d+)>')
 tag_pat=re.compile(r'<(/?)(A228401-T\d+(?:\.[^>]*)?)>')
 op_pat=re.compile(r'\[\[(/?)(MOVE|TARGET|SWAP)-(edit|convention)-V(\d+)-([\d.]+(?:-\d+)?)\]\]')
 current=None;segments={};unassigned=[];struct_events=[];code_events=[];ops=[]
 for pi,p in enumerate(doc.xpath('./w:body//w:p',namespaces=NS),1):
  # Default run formatting comes from styles, not pPr/rPr (paragraph-mark formatting).
  ps=p.find('w:pPr/w:pStyle',NS);sid=ps.get('{'+W+'}val') if ps is not None else 'Normal'
  inherited={}; chain=[];cursor=styles.get(sid)
  while cursor is not None and len(chain)<12:
   chain.append(cursor);base=cursor.find('w:basedOn',NS);cursor=styles.get(base.get('{'+W+'}val')) if base is not None else None
  for c in reversed(chain):inherited.update(attrs_rpr(c.find('w:rPr',NS)))
  runs=[];start=0
  for ri,r in enumerate(p.xpath('.//w:r',namespaces=NS),1):
   rp=attrs_rpr(r.find('w:rPr',NS));effective=dict(inherited)
   sty=r.find('w:rPr/w:rStyle',NS)
   if sty is not None and sty.get('{'+W+'}val') in styles:effective.update(attrs_rpr(styles[sty.get('{'+W+'}val')].find('w:rPr',NS)))
   effective.update(rp);pp=props(effective)
   t=rtext(r);fn=r.xpath('./w:footnoteReference/@w:id',namespaces=NS)
   raw_runs_count+=1
   c=pp.get('color');h=pp.get('highlight');k=[]
   if t:
    colors[str(c)]+=1;highlights[str(h)]+=1
    if c in ('538135','009E47'):k.append('hand')
    if c in ('C55A11','C45911'):k.append('dram_del')
    if c=='00B0F0':k+=['hand','within_deleted']
    if c=='808080':k.append('editorial')
    if c=='C00000':k.append('music')
    if pp.get('strike') not in (None,False,'false','0','off'):k.append('text_del')
    if h=='green':k.append('stage')
    if h=='cyan':k.append('heading')
    if h=='yellow':k.append('entity')
    if h=='magenta':k.append('page_marker')
    for y in k:sems[y]+=1
   runs.append({'n':ri,'a':start,'b':start+len(t),'text':t,'properties':pp,'direct_properties':props(rp),'semantics':k,'footnotes':fn})
   for f in fn:footrefs.append({'id':f,'p':pi,'offset':start,'t':current})
   start+=len(t)
  text=''.join(r['text'] for r in runs)
  p_rec={'id':f'xml-p{pi:04d}','n':pi,'paraId':p.get('{http://schemas.microsoft.com/office/word/2010/wordml}paraId'),'text':text,'runs':runs,'t':current,'role_refs':[],'footnotes':[f for r in runs for f in r['footnotes']], 'xml_path':f'/word/document.xml/body/p[{pi}]'}
  ev=list(tag_pat.finditer(text))
  for m in ev:
   closing,tid=m.groups();struct_events.append({'p':pi,'a':m.start(),'b':m.end(),'close':bool(closing),'id':tid,'raw':m.group()})
   if not closing:
    if current:issue('nested-t-'+str(pi),'error','פתיחת יחידת T בתוך יחידה פתוחה',current+' -> '+tid,{'xml_p':pi},'טעון בדיקת גבולות לפני חלוקת הטקסט.')
    if tid in segments:issue('duplicate-t-'+tid,'error','מזהה T חוזר',tid,{'xml_p':pi},'אין לאחד שתי יחידות בשקט.')
    current=tid;segments[tid]={'id':tid,'start':pi,'end':None,'paragraphs':[]}
   else:
    if current!=tid:issue('close-t-'+str(pi),'error','סגירת T שאינה תואמת לפתיחה',f'{current} / {tid}',{'xml_p':pi},'טעון בדיקה.')
    if tid in segments:segments[tid]['end']=pi
   p_rec['t']=current or tid
  for m in role_pat.finditer(text):p_rec['role_refs'].append({'id':m[1],'a':m.start(),'b':m.end()})
  for m in op_pat.finditer(text):
   closing,kind,reason,version,key=m.groups()
   ops.append({'p':pi,'t':p_rec['t'],'a':m.start(),'b':m.end(),'raw':m.group(),'close':bool(closing),'kind':kind,'reason':reason,'version':version,'key':key})
  if current in segments:segments[current]['paragraphs'].append(p_rec['id'])
  elif text.strip():unassigned.append(p_rec['id'])
  paragraphs.append(p_rec)
  for m in ev:
   if m.group(1):current=None
 if current:issue('open-t-final','error','יחידת T ללא סגירה',current,{},'לא לשער סגירה.')
 # Residual text outside structural T tags is retained.
 residuals=[]
 for pe in paragraphs:
  st=tag_pat.sub('',pe['text'])
  if tag_pat.search(pe['text']) and st.strip():residuals.append({'p':pe['n'],'t':pe['t'],'text':st,'raw':pe['text']})
 if residuals:issue('residual-struct','warning','טקסט נוסף בשורה של קוד מבנה',str(residuals),{'xml_p':residuals[0]['p']},'הטקסט נשמר כחריגה גלויה; אין למחוק אותו בעת חילוץ היחידה.')

 # Validate spans and pairings; do not silently repair mismatched closes.
 active={};spans=[];op_errors=[]
 for op in ops:
  key=(op['kind'],op['reason'],op['version'],op['key'])
  if op['kind']=='TARGET':continue
  if not op['close']:
   if key in active:op_errors.append({'type':'nested_open','event':op})
   active[key]=op
  elif key not in active:op_errors.append({'type':'unmatched_close','event':op})
  else:spans.append({'open':active.pop(key),'close':op})
 for op in active.values():op_errors.append({'type':'unclosed_open','event':op})
 movements=[]
 for span in spans:
  if span['open']['kind']=='MOVE':
   target=[o for o in ops if o['kind']=='TARGET' and all(o[k]==span['open'][k] for k in ('reason','version','key'))]
   movements.append({'id':span['open']['key'],'kind':'MOVE','valid':len(target)==1,'source':span,'target':target[0] if len(target)==1 else target})
 swaps=defaultdict(list)
 for o in ops:
  if o['kind']=='SWAP':swaps[(o['reason'],o['version'],o['key'].split('.')[0])].append(o)
 for (reason,version,page),evts in swaps.items():
  sp=[s for s in spans if s['open']['kind']=='SWAP' and s['open'] in evts]
  good=len(sp)==2 and not any(er['event'] in evts for er in op_errors)
  movements.append({'id':page,'kind':'SWAP','valid':good,'spans':sp,'events':evts})
  if not good:issue('swap-'+page,'error','זוג SWAP שלא ניתן להפעיל בבטחה','סימון פתיחה .1-2 נסגר עם .2-1; המקור אינו מתוקן אוטומטית.',{'xml_p':evts[0]['p'],'t':evts[0]['t'],'events':evts},'החילוף מוצג אך אינו מופעל במצב הערוך.')

 roles_in_xml=Counter(r['id'] for p in paragraphs for r in p['role_refs'])
 units_in_xml=set(segments); units_in_md=set(text_records)
 for t in sorted(units_in_xml-units_in_md):issue('xml-only-'+t,'error','T במקור ללא רשומה במיפתוח',t,{},'יחידה נשמרת ללא S משוער.')
 for t in sorted(units_in_md-units_in_xml):issue('md-only-'+t,'warning','T במיפתוח שלא נמצאה במקור',t,{},'לא להמציא טקסט.')
 role_cross=[]
 for p in paragraphs:
  expected=text_records.get(p['t'],{}).get('target')
  for r in p['role_refs']:
   rec=roles.get(r['id'])
   if not rec:issue('unknown-'+r['id'],'warning','קוד R ללא רשומת תפקיד',r['id'],{'xml_p':p['n']},'הקוד נשמר כקישור לא פתור.')
   elif rec['s']!=expected:role_cross.append({'p':p['n'],'t':p['t'],'role':r['id'],'t_s':expected,'r_s':rec['s']})
 if role_cross:issue('role-scene-cross','warning','הקשר תפקיד־יחידה אינו זהה בכל הופעה',f'{len(role_cross)} הופעות',{'examples':role_cross},'ייתכן אזכור/הוראת בימוי חוצה יחידות; נדרש בירור לפני הצגת קשר ודאי.')

 # Entity-highlight spans, grouped across Word runs. Name matching remains separately labelled.
 entities=[]
 for p in paragraphs:
  chunks=[];a=None;content=''
  for rr in p['runs']+[{'semantics':[],'text':'','a':len(p['text']),'b':len(p['text'])}]:
   if 'entity' in rr['semantics']:
    if a is None:a=rr['a']
    content+=rr['text']
   elif a is not None:
    name=norm(content).strip('<> ');pids=p_by_name.get(name,[])
    if re.fullmatch(r'[/\s]+',name):
     a=None;content='';continue
    entities.append({'p':p['n'],'t':p['t'],'a':a,'b':a+len(content),'raw':content,'name':name,'person':pids[0] if len(pids)==1 else None,'resolution':'unique_exact_name' if len(pids)==1 else 'unresolved'})
    used_persons.update(pids if len(pids)==1 else []);a=None;content=''

 # Footnotes have separate content, not body-text substitutions.
 missing_fn=sorted(set(f['id'] for f in footrefs)-set(footnotes))
 unref_fn=sorted(set(footnotes)-set(f['id'] for f in footrefs),key=int)
 fn_order={}
 for i,f in enumerate(footrefs,1):
  fn_order.setdefault(f['id'],len(fn_order)+1)
 for fid,f in footnotes.items():f['display_number']=fn_order.get(fid);f['occurrences']=[o for o in footrefs if o['id']==fid]
 if missing_fn:issue('fn-missing','error','הפניות להערות ללא תוכן',', '.join(missing_fn),{},'להציג הערה חסרה ולא להסיט מספור.')
 if unref_fn:issue('fn-unreferenced','warning','הערות שאינן מופנות מגוף הטקסט',str(unref_fn),{},'נשמרות בקובץ הנתונים; אין לשתול אותן בגוף.')
 comments=[]
 if '/word/comments.xml' in parts:
  for co in parts['/word/comments.xml'].xpath('.//w:comment',namespaces=NS):
   comment_id=co.get('{'+W+'}id')
   anchors=[pe['n'] for pe,element in zip(paragraphs,doc.xpath('./w:body//w:p',namespaces=NS)) if comment_id in element.xpath('.//w:commentRangeStart/@w:id | .//w:commentReference/@w:id',namespaces=NS)]
   comments.append({'id':comment_id,'paragraphs':anchors,'text':'\n'.join(''.join(p.xpath('.//w:t/text()',namespaces=NS)) for p in co.xpath('./w:p',namespaces=NS)),'kind':'work_comment'})

 # Compare page assignments without silently normalizing source values.
 page_records=matches('מספרי עמודים')
 page_keys=defaultdict(list)
 for r in page_records:page_keys[(get(r['values'],0),get(r['values'],1))].append(r['row'])
 duplicates=[{'t':key[0],'page':key[1],'rows':rr} for key,rr in page_keys.items() if len(rr)>1]
 if duplicates:issue('duplicate-page-link','warning','קישור כפול של מקטע לעמוד',str(duplicates),{'sheet':'מספרי עמודים'},'מפת הכיסוי סופרת את העמוד פעם אחת, אך שתי הרשומות נשמרות.')
 page_diffs=[]
 for tid,t in text_records.items():
  listed={int(n) for n in re.findall(r'\d+',str(get(t['record']['values'],2)).split('.')[0] if isinstance(get(t['record']['values'],2),float) else str(get(t['record']['values'],2)))}
  tab={int(r['values'][1]) for r in page_records if r['values'][0]==tid and isinstance(r['values'][1],(float,int))}
  xml_pages={int(re.search(r'<(\d+)>',p['text'])[1]) for p in paragraphs if p['t']==tid and any('page_marker' in r['semantics'] for r in p['runs']) and re.search(r'<(\d+)>',p['text'])}
  if listed!=tab or listed!=xml_pages:page_diffs.append({'t':tid,'T_pages':sorted(listed),'page_sheet':sorted(tab),'XML_markers':sorted(xml_pages)})
 if page_diffs:issue('page-assignment-mismatch','warning','אי-התאמה בין מיפויי עמודים',str(page_diffs),{'sheet':'מספרי עמודים','comparisons':page_diffs},'להציג את שלוש מערכות המיקום בנפרד, בלי לייחס מספר עמוד שהסקנו למקור.')
 # Exact joins and documented derivations (not all an ontology of claims).
 rel=[]
 def edge(source,target,kind,record,field,basis='explicit_id'):
  rel.append(dict(source=source,target=target,kind=kind,basis=basis,evidence={'sheet':record.get('sheet'),'row':record.get('row'),'cell':field+str(record.get('row',''))}))
 for tid,t in text_records.items():edge(tid,t['target'],'text_assignment',t['record'],'F')
 for pid,p in graphics.items():edge(pid,A,'belongs_to_production',p['record'],'A')
 for pid,p in progs.items():
  edge(pid,p['target'],'program_assignment',p['record'],'E')
  edge(pid,p['parent'],'program_part_of',p['record'],'D','identifier_derivation')
 for rid,r in roles.items():
  edge(rid,r['s'],'role_in_unit',r['record'],'C')
  if r['person'] in all_persons:edge(rid,r['person'],'performed_by',r['record'],'H')
 for av in media:edge(av['id'],av['target'],'media_assignment',av['record'],'D')
 for q in name_links:rel.append(dict(source=q['source'],target=q['person'],kind='named_credit',basis=q['kind'],evidence={'cell':q['source'],'raw':q['name']}))

 # Known scope/quality issues, with exact original values preserved.
 for s in sheets:
  for er in s['errors']:
   if er['target']:issue('cell-'+s['name']+er['cell'],'warning','תא שגיאה במיפתוח',er['value'],{'sheet':s['name'],'cell':er['cell']},'התא מוצג כשגיאה; אינו הופך לערך היסטורי.')
 issue('front-target','info','T1 מפנה להצגה ולא ליחידת S','A228401-T1 -> A228401',text_records[A+'-T1']['record'],'להבחין בשער כללי; זה אינו בהכרח קישור שבור.')
 issue('ps-numbering','info','מספור פנימי שונה מן המספור המודפס','11 רשומות PS הכוללות שתי הפסקות לעומת תשעה קטעים ממוספרים.',{},'אין להשתמש במספר הסצינה הפנימי כמספר המודפס בתכנייה.')
 issue('inventory-program','question','מספר פריטי תכנייה לעומת גרסאות ממופתחות','בתכולת תיקים: 2 פריטים, 4 עמודים; ב-P נמצאת רק משפחת P1 עם 6 יחידות גרפיות.',matches('תכולת תיקים')[1],'אין להסיק שקיימות שתי גרסאות תוכן שונות; נדרש לברר משמעות יחידת הספירה.')
 issue('inventory-coverage','question','מלאי חומר מול מפתוח מפורט','המלאי מונה 5 כרזות ו-3 פריטי עיתונות; בגיליונות הייעודיים אין רשומות של ההצגה.',{},'מלאי אינו זהה לנגישות קבצים או למפתוח של כל פריט.')
 issue('tender-source','question','הופעת הטנדר בטקסט לעומת הערת המיפתוח','T16 כולל "הטנדר נוסע!" וגם "טנדר". גוף השיר אינו מצוי שם. ב-T/S/PS יש הערה "לא נמצא בטקסט המודפס".',{'t':A+'-T16','xml_p':2235},'ההערה והמקור מוצגים זה לצד זה; לא טוענים שאין בטקסט כל זכר לשיר.')
 issue('program-lyrics-link','info','יחידות שיר בתכנייה ללא שדה S מאוכלס','P1.4/P1.5/P1.6 מזוהות בהערות, אך שדה הקישור H ריק.',{'sheet':choose('P -')['name'],'cells':'G246:H248'},'קשרים לפי הכותרת מוצגים כזיהוי מן ההערה, לא כמפתח זר מפורש.')
 issue('music-absence','info','לא נמצאו רשומות M של ההצגה','אין התאמה למזהי A/T/S או לשם ההצגה בגיליון M.',{'sheet':choose('M-')['name']},'זהו היעדר מפתוח, לא הוכחה שלא הייתה מוזיקה או שלא שרדו תווים.')

 issue('ps-migration-notes','warning','תוכן נלווה ב-PS אינו בהכרח חלק מן הסצנה באותה שורה','בשורות 51–52 של PS, עמודת S כוללת פזמונים מתוך פלוגה דום ובשבילך עד חצי המלכות; כותרת העמודה מורה להעביר אותם למילות השיר בתכנייה.',{'sheet':choose('PS -')['name'],'cells':'S51:S52'},'אין לייחס את הפזמונים ל-S1 ול-S2 רק משום שהם נמצאים באותן שורות; היחס צריך להיבדק מול P והצילום.')
 # Category and granular statistics for source-backed distant reading.
 segstats=[]
 for tid,s in segments.items():
  pp=[p for p in paragraphs if p['t']==tid];fs=[footnotes[f] for p in pp for f in p['footnotes'] if f in footnotes]
  cat=Counter(c for f in fs for c in f['categories']);sem=Counter(k for p in pp for r in p['runs'] if r['text'] for k in r['semantics'])
  s.update({'paragraph_count':len(pp),'nonempty_paragraphs':sum(bool(p['text'].strip()) for p in pp),'char_count':sum(len(p['text']) for p in pp),'role_codes':sum(len(p['role_refs']) for p in pp),'role_ids':sorted({r['id'] for p in pp for r in p['role_refs']}),'footnote_count':sum(len(p['footnotes']) for p in pp),'categories':dict(cat),'semantics_runs':dict(sem)})
  segstats.append({'id':tid,**{k:v for k,v in s.items() if k not in ('paragraphs','id')}})
 persons={pid:all_persons[pid] for pid in sorted(used_persons)}
 for s in sheets:
  if s['name']=='ישויות':s['indirect_rows']=len(persons);s['relevant_records']=[p['record'] for p in persons.values()]
  else:s['indirect_rows']=0;s['relevant_records']=[r for r in s['rows'] if r['direct_match']]

 stats={'sheet_count':len(sheets),'md_T':len(text_records),'xml_T':len(segments),'S':len(scene_records),'PS':len(progs),'P':len(graphics),'roles':len(roles),'referenced_role_ids':len(roles_in_xml),'role_marker_occurrences':sum(roles_in_xml.values()),'persons':len(persons),'performers_with_id':len({r['person'] for r in roles.values() if r['person'] in all_persons}),'roles_without_actor':sum(r['person'] not in all_persons for r in roles.values()),'named_entity_spans':len(entities),'xml_parts':len(parts),'paragraphs':len(paragraphs),'nonempty_paragraphs':sum(bool(p['text'].strip()) for p in paragraphs),'runs':raw_runs_count,'characters':sum(len(p['text']) for p in paragraphs),'footnotes':len(footnotes),'footnote_references':len(footrefs),'comments':len(comments),'movement_events':len(ops),'valid_MOVEs':sum(o['kind']=='MOVE' and o['valid'] for o in movements),'valid_SWAPs':sum(o['kind']=='SWAP' and o['valid'] for o in movements),'invalid_SWAPs':sum(o['kind']=='SWAP' and not o['valid'] for o in movements),'media_included':len(media),'media_quarantined':len(quarantine),'colors':dict(colors),'highlights':dict(highlights),'semantic_runs':dict(sems),'missing_footnotes':missing_fn,'unreferenced_footnotes':unref_fn,'missing_T':sorted(units_in_md-units_in_xml),'orphan_T':sorted(units_in_xml-units_in_md),'role_scene_cross':role_cross}
 model={'production':{'id':A,'title':TITLE,'record':prod},'provenance':{'spreadsheet_id':'1gVFSNSJHc1W4QMKXrXuSIdO2kuGvNH_c6mRihNuuxLQ','spreadsheet_modified':'2026-10-10T11:16:07.096Z','snapshot_date':'2026-10-12','xml_file':xmlpath.name,'xml_type':'Word Flat OPC / WordprocessingML; not TEI','sha256':{'xml':sha(xmlpath),'md':sha(xlsxpath)},'reference':'archived 10.2; the supplied GitHub Pages URL could not be fetched in this session'},'stats':stats,'sheets':[{k:v for k,v in s.items() if k!='rows'} for s in sheets],'text_records':text_records,'scenes':scene_records,'program_items':progs,'graphics':graphics,'roles':roles,'persons':persons,'persons_name_links':name_links,'media':media,'quarantine':quarantine,'paragraphs':paragraphs,'segments':segments,'footnotes':footnotes,'comments':comments,'entity_mentions':entities,'operations':movements,'raw_operation_events':ops,'relations':rel,'issues':issues,'structural_events':struct_events,'residuals':residuals,'page_comparisons':page_diffs,'duplicate_page_links':duplicates}
 jwrite(P/'data'/'edition.json',model);jwrite(P/'data'/'workbook_all_tabs.json',sheets)
 jwrite(P/'data'/'quality_report.json',{'stats':stats,'issues':issues})
 summary=P/'data'/'extraction_summary.json';jwrite(summary,{'stats':stats,'issues':issues,'entities':entities,'persons':list(persons),'segments':segstats})
 print(json.dumps(stats,ensure_ascii=False,indent=2))
 print('\nVALID operations',[(o['kind'],o['id'],o['valid']) for o in movements])
 print('\nROSTER',[(k,r['character'],r['actor'],r['person']) for k,r in roles.items()])
 return model
if __name__=='__main__':
 ap=argparse.ArgumentParser();ap.add_argument('--xml',type=Path);ap.add_argument('--md',type=Path);a=ap.parse_args()
 run(a.xml or P/'sources'/'ki_tetze_23_9_26.xml',a.md or P/'sources'/'master_MD_snapshot.xlsx')
